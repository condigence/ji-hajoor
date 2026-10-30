"""Adds a "What changed on 30 October 2026" slide as slide 2 of the two existing decks
(JiHajoor_Architecture.pptx and JiHajoor_Design_Deck.pptx), styled like each deck.

Usage, from the repository root:  python3 docs/source/update_existing_decks.py
Running it again does nothing: a deck that already has the slide is left alone.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DOCS = Path(__file__).resolve().parent.parent
MARKER = 'What changed on 30 October 2026'

CARDS = [
    ('Security review', [
        'Web sessions: refresh token in an HttpOnly, Secure, SameSite=Strict cookie',
        'Passwords: 10+ characters, letter and digit; common, sequence and personal ones refused',
        'Rate limits keyed on Railway’s X-Real-IP; https-only image links; masked query strings in logs',
        'npm audit: 0 vulnerabilities in all three apps',
    ]),
    ('Reviews engine', [
        'Review prompt after payment, one reminder after 24 h (ReviewReminder: 14 event types now)',
        'Public page per verified worker: /public/workers/{id} and its reviews',
        'Workers reply to reviews; share-my-profile card in the partner app',
        'Booking-flow CI now covers prompt, moderation, public page and reply',
    ]),
    ('Names used everywhere', [
        'Apps: Ji-Hajoor (customers), Ji-Hajoor Pro (workers), Ji-Hajoor Admin (staff, My agency)',
        'People: customer, worker (Provider in code), agency (Tenant in code)',
        'Agencies invite workers; agency share 0–50% of the worker’s net',
    ]),
    ('Read next', [
        'JiHajoor_Tech_Stack.pptx: current stack and security state',
        'JiHajoor_Pitch_Deck, Revenue_Model, Brand_and_Marketing',
        'User guides: docs/guides/; PDFs in docs/pdf/',
        'Security findings later in this deck show their state on 18 October',
    ]),
]


def has_marker(prs):
    return any(MARKER in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


def move_to(prs, slide, index):
    ids = prs.slides._sldIdLst
    el = [e for e in ids if prs.slides.part.related_part(e.rId) is slide.part][0]
    ids.remove(el)
    ids.insert(index, el)


def add_cards(slide, top, colors, card_fill, font=None):
    title_rgb, body_rgb, accent_rgb = colors
    x0, w, h, gap = 0.6, 5.95, 2.45, 0.25
    for i, (title, lines) in enumerate(CARDS):
        x = x0 + (i % 2) * (w + gap)
        y = top + (i // 2) * (h + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.adjustments[0] = 0.05
        box.fill.solid()
        box.fill.fore_color.rgb = card_fill
        box.line.fill.background()
        box.shadow.inherit = False
        tb = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.15), Inches(w - 0.5), Inches(h - 0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size, r.font.bold, r.font.color.rgb = Pt(17), True, title_rgb
        if font:
            r.font.name = font
        for line in lines:
            p = tf.add_paragraph()
            p.space_before = Pt(4)
            r = p.add_run()
            r.text = '•  ' + line
            r.font.size, r.font.color.rgb = Pt(13), body_rgb
            if font:
                r.font.name = font


def architecture():
    path = DOCS / 'JiHajoor_Architecture.pptx'
    prs = Presentation(path)
    if has_marker(prs):
        return print('skip', path.name)
    layout = prs.slides[1].slide_layout
    slide = prs.slides.add_slide(layout)
    phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    ordered = sorted(phs)
    phs[ordered[0]].text = MARKER
    if len(ordered) > 1:
        phs[ordered[1]].text = 'Security review, the reviews engine and naming: what this deck does not show yet'
    for idx in ordered[2:]:
        phs[idx]._element.getparent().remove(phs[idx]._element)
    add_cards(slide, 1.75, (RGBColor(0x1D, 0x29, 0x39), RGBColor(0x34, 0x40, 0x54), None), RGBColor(0xF2, 0xF4, 0xF7))
    move_to(prs, slide, 1)
    prs.save(path)
    print('updated', path.name)


def design():
    path = DOCS / 'JiHajoor_Design_Deck.pptx'
    prs = Presentation(path)
    if has_marker(prs):
        return print('skip', path.name)
    slide = prs.slides.add_slide(prs.slides[1].slide_layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    k = slide.shapes.add_textbox(Inches(0.89), Inches(0.62), Inches(11.5), Inches(0.3)).text_frame
    k.paragraphs[0].alignment = PP_ALIGN.LEFT
    r = k.paragraphs[0].add_run()
    r.text = 'UPDATE · 30 OCTOBER 2026'
    r.font.name, r.font.size, r.font.color.rgb = 'Courier New', Pt(12), RGBColor(36, 84, 197)
    t = slide.shapes.add_textbox(Inches(0.89), Inches(0.95), Inches(11.5), Inches(0.6)).text_frame
    t.paragraphs[0].alignment = PP_ALIGN.LEFT
    r = t.paragraphs[0].add_run()
    r.text = MARKER
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = 'Calibri', Pt(30), True, RGBColor(31, 58, 112)
    add_cards(slide, 1.75, (RGBColor(31, 58, 112), RGBColor(74, 88, 120), None), RGBColor(0xEE, 0xF2, 0xFD), 'Calibri')
    move_to(prs, slide, 1)
    prs.save(path)
    print('updated', path.name)


if __name__ == '__main__':
    architecture()
    design()
